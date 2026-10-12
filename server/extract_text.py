"""材料 → 文本。

agent 不接触用户文件（不挂盘），所以材料必须先在这里变成文本，
再由服务端拼进 prompt。这一层是**确定性代码**，按用户隔离、
跑在服务端进程里，与 agent 的权限完全无关。

支持：
    .txt .md .csv .json .log               直接读（UTF-8 / GBK 兜底）
    .docx                                  python-docx（段落 + 表格）
    .pdf                                   pypdf 逐页
    .pptx                                  python-pptx
    .xlsx                                  openpyxl（限行数）
    .doc .rtf .odt .wps 等                 交给 LibreOffice 转 txt 兜底
    图片                                  不做 OCR，明确返回「未提取到文本」

任何一步失败都不抛错：返回 parse_error，材料照样入库（用户能在界面看到原因）。
"""

from __future__ import annotations

import json
import subprocess
import tempfile
from pathlib import Path

# 单个材料最多取多少字符：再多也塞不进上下文，截断并如实标注
MAX_CHARS = 60_000
# 表格/表格型文件的取用上限
MAX_TABLE_ROWS = 300
MAX_SHEET_CELLS = 6_000

PLAIN_SUFFIXES = {".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".log", ".yaml", ".yml", ".xml", ".html", ".htm"}
LO_SUFFIXES = {".doc", ".rtf", ".odt", ".wps", ".pages", ".dot", ".dotx"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff", ".heic"}
SOFFICE = "soffice"


def _read_plain(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8", "utf-8-sig", "gb18030", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def _read_docx(path: Path) -> str:
    from docx import Document  # 延迟导入：没装也不影响其他格式

    doc = Document(str(path))
    parts = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    for index, table in enumerate(doc.tables[:20], 1):
        parts.append(f"[表 {index}]")
        for row in table.rows[:MAX_TABLE_ROWS]:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            parts.append(" | ".join(cells))
    return "\n".join(parts)


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for number, page in enumerate(reader.pages, 1):
        try:
            text = page.extract_text() or ""
        except Exception:  # noqa: BLE001 单页坏页不该让整份材料失败
            text = ""
        if text.strip():
            pages.append(f"----- 第 {number} 页 -----\n{text.strip()}")
    return "\n\n".join(pages)


def _read_pptx(path: Path) -> str:
    from pptx import Presentation

    deck = Presentation(str(path))
    slides = []
    for number, slide in enumerate(deck.slides, 1):
        lines = []
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text and text.strip():
                lines.append(text.strip())
        if lines:
            slides.append(f"----- 第 {number} 页 -----\n" + "\n".join(lines))
    return "\n\n".join(slides)


def _read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    book = load_workbook(str(path), read_only=True, data_only=True)
    lines: list[str] = []
    budget = MAX_SHEET_CELLS
    for sheet in book.worksheets:
        lines.append(f"----- 工作表：{sheet.title} -----")
        for row in sheet.iter_rows(values_only=True, max_row=MAX_TABLE_ROWS):
            cells = ["" if cell is None else str(cell) for cell in row]
            if not any(c.strip() for c in cells):
                continue
            lines.append(" | ".join(cells))
            budget -= len(cells)
            if budget <= 0:
                break
        if budget <= 0:
            break
    book.close()
    return "\n".join(lines)


def _convert_with_libreoffice(path: Path, workdir: Path) -> str:
    """老格式（.doc/.rtf/.odt…）交给容器里已有的 LibreOffice 转成 txt。"""
    result = subprocess.run(
        [SOFFICE, "--headless", "--norestore", "--convert-to", "txt:Text",
         "--outdir", str(workdir), str(path)],
        capture_output=True, text=True, timeout=180,
    )
    if result.returncode != 0:
        raise RuntimeError(f"LibreOffice 转换失败：{result.stderr.strip()[:200]}")
    produced = workdir / (path.stem + ".txt")
    if not produced.exists():
        raise RuntimeError("LibreOffice 没有产出文本（格式不受支持？）")
    return _read_plain(produced)


def extract_text(path: Path, *, max_chars: int = MAX_CHARS) -> dict:
    """把一份材料变成文本。永不抛错，失败信息放 parse_error。"""
    suffix = path.suffix.lower()
    name = path.name
    size = path.stat().st_size if path.exists() else 0
    text = ""
    parse_error: str | None = None

    try:
        if suffix in IMAGE_SUFFIXES:
            parse_error = "图片材料未做 OCR：本版只把文件名与用途告诉 agent，正文请补文字说明"
        elif suffix in PLAIN_SUFFIXES:
            text = _read_plain(path)
        elif suffix == ".docx":
            text = _read_docx(path)
        elif suffix == ".pdf":
            text = _read_pdf(path)
        elif suffix == ".pptx":
            text = _read_pptx(path)
        elif suffix == ".xlsx":
            text = _read_xlsx(path)
        elif suffix in LO_SUFFIXES:
            with tempfile.TemporaryDirectory() as tmp:
                text = _convert_with_libreoffice(path, Path(tmp))
        else:
            # 认不出后缀：先当文本试试，不行就交给 LibreOffice
            try:
                text = _read_plain(path)
                if "\x00" in text[:1000]:
                    raise ValueError("binary")
            except Exception:  # noqa: BLE001
                with tempfile.TemporaryDirectory() as tmp:
                    text = _convert_with_libreoffice(path, Path(tmp))
    except ModuleNotFoundError as exc:
        parse_error = f"服务端缺少解析依赖（{exc.name}），已跳过正文提取"
    except Exception as exc:  # noqa: BLE001 任何解析失败都不该挡住上传
        parse_error = f"解析失败：{type(exc).__name__}: {str(exc)[:160]}"

    text = (text or "").strip()
    truncated = False
    if len(text) > max_chars:
        text = text[:max_chars]
        truncated = True
    if not text and not parse_error:
        parse_error = "未提取到文本（可能是扫描件或纯图片文档）"

    return {
        "name": name,
        "bytes": size,
        "chars": len(text),
        "truncated": truncated,
        "text": text,
        "parse_error": parse_error,
    }


def dumps(meta: dict) -> str:
    """材料索引（不含正文）落盘用；正文单独存在大字段里。"""
    return json.dumps(
        {k: v for k, v in meta.items() if k != "text"}, ensure_ascii=False, indent=2
    )
