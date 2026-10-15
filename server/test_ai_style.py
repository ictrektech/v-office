"""整篇样式指令（<doc-style … />）的链路测试。

这个功能有一类特别难查的失败：**代码跑了，但没体现**——属性名解析错、参数没传到
脚本、值被静默丢掉，界面上看起来都只是「AI 说改好了，字号却没变」。所以这里不只测
解析，而是一路测到**导出的 docx**：真跑一次 md_to_docx.py，再把 run 的字号读回来。

跑法：
    cd server && python3 -m unittest test_ai_style -v
"""

import logging
import os
import tempfile
import unittest
from pathlib import Path

# AI_HOME 是模块级常量（import 时就定了），要在 import ai **之前**指到临时目录，
# 否则测试会往真实的 /data/.ai 里写草稿和会话。
_TMP_HOME = tempfile.mkdtemp(prefix="vo-ai-style-home-")
os.environ.setdefault("V_OFFICE_AI_HOME", _TMP_HOME)

import ai  # noqa: E402
import docx_revise  # noqa: E402

# ai.LOG 默认 None（由 main.py 启动时注入）。不接上，转换一出错就是
# `'NoneType' has no attribute 'warning'`，真实失败原因反而被盖掉。
ai.LOG = logging.getLogger("test-ai-style")


class DocStyleParseTest(unittest.TestCase):
    """<doc-style/> 的解析与剥离。"""

    def test_style_and_text_patch_coexist(self):
        out = ai.parse_model_output(
            '这段改一下'
            '<doc-style body-size="14" body-font="宋体" line-spacing="1.5"/>'
            '<patch anchor="2">新文本</patch>'
        )
        self.assertEqual(
            out["doc_style"],
            {"font": "宋体", "size_pt": 14.0, "line_spacing": 1.5},
        )
        self.assertEqual([p["anchor"] for p in out["patches"]], ["2"])
        self.assertEqual(out["prose"], "这段改一下")

    def test_style_tag_never_leaks_into_prose(self):
        """样式标签绝不能当正文写进草稿（以前 <final-doc> 就这样漏过）。"""
        out = ai.parse_model_output('<doc-style body-size="12"/>这是说明')
        self.assertNotIn("doc-style", out["prose"])
        self.assertNotIn("body-size", out["prose"])
        self.assertEqual(out["prose"], "这是说明")

    def test_style_only_turn(self):
        out = ai.parse_model_output('<doc-style heading-size="16"/>')
        self.assertEqual(out["doc_style"], {"heading_size_pt": 16.0})
        self.assertEqual(out["patches"], [])
        self.assertEqual(out["prose"], "")

    def test_underscore_and_hyphen_both_accepted(self):
        # 曾经踩过：属性正则的 \w+ 匹配不了连字符，body-size 被读成 size，整条指令静默失效
        hyphen = ai.parse_model_output('<doc-style body-size="14"/>')["doc_style"]
        under = ai.parse_model_output('<doc-style body_size="14"/>')["doc_style"]
        self.assertEqual(hyphen, {"size_pt": 14.0})
        self.assertEqual(under, {"size_pt": 14.0})

    def test_out_of_range_values_rejected(self):
        """这些值会直接写进 docx——字号 9999 会让用户拿到一份没法看的文档。"""
        out = ai.parse_model_output(
            '<doc-style body-size="9999" heading-size="0" line-spacing="12"/>'
        )
        self.assertEqual(out["doc_style"], {})

    def test_plain_output_unaffected(self):
        out = ai.parse_model_output("就一段普通的说明文字")
        self.assertEqual(out["doc_style"], {})
        self.assertEqual(out["prose"], "就一段普通的说明文字")


class StyleArgsTest(unittest.TestCase):
    """草稿 styles → md_to_docx.py 命令行参数。"""

    def test_builds_expected_argv(self):
        args = ai._style_args(
            {"styles": {"font": "宋体", "size_pt": 14.0, "line_spacing": 1.5}}
        )
        self.assertEqual(
            args, ["--body-font", "宋体", "--body-size", "14.0", "--line-spacing", "1.5"]
        )

    def test_no_styles_means_no_args(self):
        """没样式的稿子，转换命令必须与从前逐字一致——这功能对既有链路是零影响。"""
        self.assertEqual(ai._style_args({}), [])
        self.assertEqual(ai._style_args({"styles": {}}), [])
        self.assertEqual(ai._style_args({"styles": {"font": "", "size_pt": None}}), [])


class StyleConversionTest(unittest.IsolatedAsyncioTestCase):
    """端到端：样式真的落到 docx 的 run 上（不只是"参数传出去了"）。"""

    MD = "# 测试标题\n\n这是正文段落，用来验证字号。\n\n行内代码 `code_here` 不该被改。\n"

    async def _convert(self, styles: dict, prefer: str = "") -> bytes:
        with tempfile.TemporaryDirectory() as tmp:
            data = await ai._convert_docx(
                tool="md_to_docx.py",
                src=self.MD.encode("utf-8"),
                src_suffix=".md",
                out_suffix=".docx",
                root=Path(tmp),
                prefer=prefer,
                extra_args=ai._style_args({"styles": styles}),
            )
        return data

    async def test_every_skill_pack_supports_style(self):
        """技能包里**每一份** md_to_docx.py 都得认样式参数。

        `_skill_tool` 优先用「本次会话所属技能包」的那份同名脚本，所以只改一份的
        后果很隐蔽：用另一个技能写稿时，脚本不认 --body-size、转换直接失败，
        用户看到的是「AI 说改好了，文档却没更新」。
        """
        from docx import Document
        import io

        packs = sorted(
            p.parent.parent.name for p in Path(ai.SKILLS_DIR).glob("*/tools/md_to_docx.py")
        )
        self.assertTrue(packs, "技能目录下找不到 md_to_docx.py")
        for pack in packs:
            with self.subTest(pack=pack):
                data = await self._convert({"size_pt": 14.0}, prefer=pack)
                self.assertIsNotNone(data, f"{pack} 的转换脚本不支持 --body-size")
                doc = Document(io.BytesIO(data))
                sizes = [
                    r.font.size.pt
                    for p in doc.paragraphs
                    for r in p.runs
                    if "这是正文" in r.text and r.font.size
                ]
                self.assertEqual(sizes[:1], [14.0], f"{pack} 的字号没生效")

    def setUp(self):
        if ai._skill_tool("md_to_docx.py", "") is None:
            self.skipTest("技能包未就绪（md_to_docx.py 不在 SKILLS_DIR），跳过转换测试")

    async def test_body_size_reaches_docx(self):
        from docx import Document
        import io

        data = await self._convert({"size_pt": 14.0})
        self.assertIsNotNone(data, "转换失败：拿不到 docx 字节")
        doc = Document(io.BytesIO(data))
        sizes = {
            r.text[:8]: (r.font.size.pt if r.font.size else None)
            for p in doc.paragraphs
            for r in p.runs
            if r.text.strip()
        }
        body = [v for k, v in sizes.items() if "这是正文" in k]
        self.assertTrue(body and body[0] == 14.0, f"正文字号没生效：{sizes}")

    async def test_mono_run_not_touched(self):
        """等宽 run 是代码，统一改样式时必须跳过，否则代码跟着正文一起变样。"""
        from docx import Document
        import io

        data = await self._convert({"size_pt": 14.0})
        doc = Document(io.BytesIO(data))
        mono = [
            r
            for p in doc.paragraphs
            for r in p.runs
            if (r.font.name or "") == "Consolas"
        ]
        self.assertTrue(mono, "找不到等宽 run，测试样本失效")
        for r in mono:
            self.assertEqual(r.font.size.pt, 9.0, "等宽 run 被误改了字号")

    async def test_default_conversion_unchanged(self):
        """不传样式时，正文仍是原来的 10.5：确认新参数默认完全不改变旧行为。"""
        from docx import Document
        import io

        data = await self._convert({})
        doc = Document(io.BytesIO(data))
        body = [
            r.font.size.pt
            for p in doc.paragraphs
            for r in p.runs
            if "这是正文" in r.text and r.font.size
        ]
        self.assertEqual(body[:1], [10.5])


class StyleSummaryTest(unittest.TestCase):
    """_style_summary：落地后那句要显示给用户的「人话」。"""

    def test_lists_fields_in_natural_order(self):
        # 字段顺序按「正文字体 → 正文字号 → 行距」，与用户下指令的顺序一致，
        # 不按 dict 的插入顺序——模型给的属性顺序不固定。
        s = ai._style_summary({"line_spacing": 1.15, "size_pt": 14.0, "font": "宋体"})
        self.assertEqual(s, "整篇样式已更新：正文字体 宋体、正文字号 14 磅、行距 1.15 倍")

    def test_float_tail_dropped(self):
        """14.0 说成「14 磅」而不是「14.0 磅」；1.15 这类真小数保持原样。"""
        self.assertIn("正文字号 14 磅", ai._style_summary({"size_pt": 14.0}))
        self.assertIn("行距 1.15 倍", ai._style_summary({"line_spacing": 1.15}))

    def test_empty_does_not_crash(self):
        self.assertEqual(ai._style_summary({}), "整篇样式已更新")


class StyleLandTurnTest(unittest.IsolatedAsyncioTestCase):
    """样式回合**真的**发得出 style 事件。

    这层曾经整个漏掉：`_land_turn` 调用了一个从未定义的 `_style_summary`，于是模型
    每次输出 <doc-style/> 都在这里抛 NameError，用户看到的是「改稿落地失败:
    name '_style_summary' is not defined」——而外面三层测试（解析 / 参数 / 真导出
    docx）全绿，谁也没发现，因为它们都不经过事件分支。样式却已经存进草稿了，
    于是症状变成「字号其实改了，但整轮报错」，比单纯失败更难查。

    autosave=False 跳过落盘回调：这个测试只管事件，不该依赖 save_document 注入。
    """

    async def test_style_only_turn_emits_style_event(self):
        uid, sid = "tester", "style-land-turn"
        ai.save_draft(uid, sid, {
            "title": "测试稿",
            "blocks": [{"heading": "一", "text": "正文段落"}],
        })
        parsed = ai.parse_model_output(
            '<doc-style body-font="宋体" body-size="14" line-spacing="1.15"/>'
        )
        events = await ai._land_turn(uid, sid, parsed, save=True, autosave=False)

        styles = [e for e in events if e.get("type") == "style"]
        self.assertEqual(
            len(styles), 1, f"没有 style 事件，实际：{[e.get('type') for e in events]}"
        )
        self.assertEqual(
            styles[0]["styles"],
            {"font": "宋体", "size_pt": 14.0, "line_spacing": 1.15},
        )
        self.assertIn("正文字号 14 磅", styles[0]["message"])
        # 样式要真落进草稿：下一轮导出 docx 靠的就是它
        self.assertEqual(
            (ai.load_draft(uid, sid).get("styles") or {}).get("size_pt"), 14.0
        )

    async def test_identical_style_reports_nothing(self):
        """同样的样式再来一次不该反复报「已更新」——用户只会觉得 AI 在空转。"""
        uid, sid = "tester", "style-land-turn-idempotent"
        ai.save_draft(uid, sid, {
            "title": "测试稿",
            "blocks": [{"heading": "一", "text": "正文段落"}],
            "styles": {"size_pt": 14.0},
        })
        parsed = ai.parse_model_output('<doc-style body-size="14"/>')
        events = await ai._land_turn(uid, sid, parsed, save=True, autosave=False)
        self.assertEqual([e for e in events if e.get("type") == "style"], [])


class ScopedStyleTest(unittest.IsolatedAsyncioTestCase):
    """**按段**样式（<doc-style anchor="N"/>）：

    用户在编辑器里选中一段说「字体变大」，结果**整篇**被一起放大——根因是
    样式只有整篇一个口径，模型输出的 <doc-style/> 不带作用域，落地时统一套到
    全文。这组测试锁住新口径：带 anchor 只改那一段、用户选区优先、真 docx 上
    只有选中段的 run 变样。
    """

    def setUp(self):
        if ai._skill_tool("md_to_docx.py", "") is None:
            self.skipTest("技能包未就绪（md_to_docx.py 不在 SKILLS_DIR），跳过转换测试")

    async def _base_docx(self) -> bytes:
        md = "## 一\n\n第一段正文。\n\n## 二\n\n第二段正文。\n"
        with tempfile.TemporaryDirectory() as tmp:
            data = await ai._convert_docx(
                tool="md_to_docx.py", src=md.encode("utf-8"),
                src_suffix=".md", out_suffix=".docx", root=Path(tmp), prefer="",
            )
        return data

    def test_parse_anchor_scope(self):
        out = ai.parse_model_output('<doc-style anchor="2" body-size="16"/>')
        self.assertEqual(
            out["doc_style"], {"size_pt": 16.0, "_anchor": "2"},
            "anchor 是作用域，不是样式字段，但必须被解析出来",
        )

    async def test_scoped_land_turn_never_touches_whole_doc_styles(self):
        """按段样式**不进**草稿 styles：styles 是整篇口径，混进去下次整篇导出
        会把选区字号铺到全文——等于换一种方式复发原 bug。"""
        uid, sid = "tester", "scoped-style-land"
        ai.save_draft(uid, sid, {
            "title": "测试稿",
            "blocks": [{"heading": "一", "text": "第一段"}, {"heading": "二", "text": "第二段"}],
        })
        parsed = ai.parse_model_output('<doc-style anchor="2" body-size="18"/>')
        events = await ai._land_turn(
            uid, sid, parsed, save=True, autosave=False, anchor="2"
        )
        styles = [e for e in events if e.get("type") == "style"]
        self.assertEqual(len(styles), 1)
        self.assertEqual(styles[0].get("scope"), "selection")
        self.assertEqual(styles[0].get("anchor"), "2")
        self.assertNotIn("整篇", styles[0]["message"])
        # 草稿 styles 必须纹丝不动
        self.assertEqual(ai.load_draft(uid, sid).get("styles") or {}, {})

    async def test_selection_anchor_beats_model_anchor(self):
        """模型把段号猜错时，以用户的选区为准（与 patch 锚点校正同一口径）。"""
        uid, sid = "tester", "scoped-style-anchor-fix"
        ai.save_draft(uid, sid, {
            "title": "测试稿",
            "blocks": [{"heading": "一", "text": "第一段"}, {"heading": "二", "text": "第二段"}],
        })
        parsed = ai.parse_model_output('<doc-style anchor="1" body-size="18"/>')
        events = await ai._land_turn(
            uid, sid, parsed, save=True, autosave=False, anchor="2"
        )
        styles = [e for e in events if e.get("type") == "style"]
        self.assertEqual(len(styles), 1)
        self.assertEqual(styles[0].get("anchor"), "2")

    async def test_only_target_paragraph_grows(self):
        """真 docx 上验证：只有选中段变大，其余段一字不动。"""
        from docx import Document
        import io

        base = await self._base_docx()
        data, applied = __import__("docx_revise").apply_scoped_style(
            base, {"size_pt": 18.0}, before="第二段正文。", index=0
        )
        self.assertEqual(applied, 1)
        before = Document(io.BytesIO(base))
        after = Document(io.BytesIO(data))
        def sizes(doc):
            return {
                p.text: [r.font.size.pt for r in p.runs if r.font.size]
                for p in doc.paragraphs
            }
        target = [k for k in sizes(after) if "第二段" in k]
        other = [k for k in sizes(after) if "第一段" in k]
        self.assertTrue(target and other, f"段落结构变了：{list(sizes(after))}")
        self.assertEqual(
            sizes(after)[target[0]], [18.0], "选中的段落没变大"
        )
        self.assertEqual(
            sizes(before)[other[0]], sizes(after)[other[0]],
            "没选中的段落也被改了——这就是用户报的全篇变大",
        )

    async def test_miss_leaves_doc_untouched(self):
        """定位不到段落就返回原文：改错段落比不改更糟。"""
        base = await self._base_docx()
        data, applied = __import__("docx_revise").apply_scoped_style(
            base, {"size_pt": 18.0}, before="不存在的段落", index=99
        )
        self.assertEqual(applied, 0)
        self.assertEqual(data, base)


class TurnPromptSelectionScopeTest(unittest.TestCase):
    """选区回合的 prompt：走**文字级**契约，不再让模型猜段落号。

    背景：用户已经选中了一段文字，服务端**直接用这段文字当坐标**去替换最稳；
    「选区 → 段号」的换算在编辑器（Writer 段落）与草稿（Markdown 按标题切块）
    划分不一致时会失效——没有标题的稿子整篇就是一块，选区作用域直接退化成
    「全篇」。所以主契约是 `<replace-selection>`，段号只留作兜底提示。
    """

    def _prompt(self, sid: str, *, anchor: str = "", selection: str = "") -> str:
        ai.save_draft("tester", sid, {
            "title": "测试稿",
            "blocks": [
                {"heading": "一", "text": "第一段正文"},
                {"heading": "二", "text": "第二段正文"},
                {"heading": "三", "text": "第三段正文"},
            ],
        })
        return ai._build_turn_prompt(
            "tester", sid, "把字体变大，这段中文改成英文",
            anchor=anchor, selection=selection,
        )

    def test_scoped_turn_uses_text_level_contract(self):
        """有选区就给文字级契约：交回替换结果，明确禁止 patch / final-doc。"""
        prompt = self._prompt("prompt-scope-01", selection="第一段正文\n第二段正文")
        self.assertIn("【只改选中的这段文字】", prompt)
        self.assertIn("<replace-selection>替换后的文字</replace-selection>", prompt)
        self.assertIn("不要用 <patch>", prompt)
        # 段号仍作为兜底信息交代（模型实在要按段改时用）
        self.assertIn("选区覆盖的段落编号：1、2", prompt)
        self.assertIn("只允许第 1、2 段", prompt)
        # 草稿里把边界标出来
        self.assertIn("[1] 一 ← 本次只改这一段", prompt)
        self.assertNotIn("[3] 三 ← 本次只改这一段", prompt)

    def test_single_paragraph_selection_also_text_level(self):
        prompt = self._prompt("prompt-scope-02", anchor="2", selection="第二段正文")
        self.assertIn("段落编号：2", prompt)
        self.assertIn("【只改选中的这段文字】", prompt)
        self.assertIn("只允许第 2 段", prompt)

    def test_selection_not_locatable_still_uses_contract(self):
        """选区与草稿对不上：点不出段号，但文字级契约照旧——这正是它的长处。"""
        prompt = self._prompt("prompt-scope-03", selection="草稿里根本没有的这句话")
        self.assertIn("【只改选中的这段文字】", prompt)
        self.assertIn("<replace-selection>", prompt)
        # 没有段号就别在 prompt 里编一个出来
        self.assertNotIn("选区覆盖的段落编号：", prompt)

    def test_no_selection_has_no_scope_block(self):
        """没选区就不该出现作用域那一段（否则等于凭空收窄）。"""
        prompt = self._prompt("prompt-scope-04")
        # 查块标题而不是裸字样：OUTPUT_CONTRACT 里也引用了「本次要改的位置」
        self.assertNotIn("# 本次要改的位置", prompt)
        self.assertNotIn("【只改选中的这段文字】", prompt)
        self.assertNotIn("← 本次只改这一段", prompt)

    def test_anchor_without_selection_keeps_plain_locate(self):
        """只给段号（在草稿视图点段定位）维持原样：不加文字级契约，
        免得把「点一段、问整篇问题」也变成限制范围。"""
        prompt = self._prompt("prompt-scope-05", anchor="1")
        self.assertIn("段落编号：1", prompt)
        self.assertNotIn("【只改选中的这段文字】", prompt)


class SelectionBlocksTest(unittest.TestCase):
    """「用户选中的区域覆盖哪几段」的推算。

    这是后续两件事的地基：prompt 里告诉模型段号范围、落文时拦越界改动。
    它算错（多圈一段 / 少圈一段）比不算更糟，所以边界场景逐个锁死。
    """

    DRAFT = {
        "title": "测试稿",
        "blocks": [
            {"heading": "一", "text": "本发明公开了一种涂覆层的制备方法。"},
            {"heading": "", "text": "其中，涂覆层包含陶瓷颗粒与纤维素。"},
            {"heading": "", "text": "陶瓷颗粒的粒径为10至50微米。"},
            {"heading": "二", "text": "与现有技术相比，本发明的有益效果是成本低。"},
        ],
    }

    def test_front_anchor_wins(self):
        """前端已定位到的段号最可信：直接采信，连选区原文都不必看。"""
        self.assertEqual(ai._selection_blocks(self.DRAFT, "随手写的文字", "3"), [3])
        self.assertEqual(ai._selection_blocks(self.DRAFT, "", "4"), [4])

    def test_anchor_out_of_range_is_dropped(self):
        self.assertEqual(ai._selection_blocks(self.DRAFT, "", "99"), [])

    def test_single_paragraph_selection(self):
        self.assertEqual(ai._selection_blocks(self.DRAFT, "涂覆层包含陶瓷颗粒", ""), [2])

    def test_selection_ignores_whitespace(self):
        """换行/空格不该影响「是不是同一句」（编辑器复制出来的文本常带缩进）。"""
        self.assertEqual(
            ai._selection_blocks(self.DRAFT, "涂覆层\n包含   陶瓷颗粒", ""), [2]
        )

    def test_cross_paragraph_selection_gives_range(self):
        """跨段选区——本次报的 bug 场景，跨段必须能推出闭区间。"""
        sel = "涂覆层包含陶瓷颗粒与纤维素。\n陶瓷颗粒的粒径为10至50微米。"
        self.assertEqual(ai._selection_blocks(self.DRAFT, sel, ""), [2, 3])

    def test_cross_paragraph_from_middle_to_middle(self):
        """从一段中间拖到另一段中间：最常见的框选形态。"""
        sel = "陶瓷颗粒与纤维素。\n陶瓷颗粒的粒径"
        self.assertEqual(ai._selection_blocks(self.DRAFT, sel, ""), [2, 3])

    def test_selection_not_in_draft_yields_empty(self):
        """选区在草稿里对不上（刚手改过还没回读）：返回空表，绝不硬猜。"""
        self.assertEqual(
            ai._selection_blocks(self.DRAFT, "这句话草稿里根本没有出现过", ""), []
        )

    def test_short_selection_does_not_cross_infer(self):
        """太短的选区不做跨段推断：一个词在多段里都出现时，圈多比圈少更糟。"""
        self.assertEqual(ai._selection_blocks(self.DRAFT, "铝合金", ""), [])

    def test_empty_inputs(self):
        self.assertEqual(ai._selection_blocks(self.DRAFT, "", ""), [])
        self.assertEqual(ai._selection_blocks({"blocks": []}, "任意文字", ""), [])


class ScopedLandTurnGuardTest(unittest.IsolatedAsyncioTestCase):
    """选区回合的落文硬校验：模型越界的改动**进不了文档**。

    这是「选中一块区域说字体变大 / 中文变英文，结果全篇被改」的机制性修复——
    不靠模型自觉，靠这里逐条校验。所以每个用例都要盯住两件事：
    草稿**一字未动**（越界那些），以及**有没有明确告诉用户被拦了**
    （静默收场会变成「AI 说改好了却没生效」，那是这个功能最坑的失败模式）。
    """

    BLOCKS = [
        {"heading": "一", "text": "第一段正文"},
        {"heading": "二", "text": "第二段正文"},
        {"heading": "三", "text": "第三段正文"},
    ]

    def _draft(self, sid: str) -> None:
        ai.save_draft("tester", sid, {"title": "测试稿", "blocks": [dict(b) for b in self.BLOCKS]})

    def _texts(self, sid: str) -> list[str]:
        return [b["text"] for b in ai.load_draft("tester", sid)["blocks"]]

    async def test_all_out_of_scope_is_blocked(self):
        """模型按整篇改（patch 全落在选区外）→ 整轮拦下，草稿一字不动。"""
        sid = "guard-all-out"
        self._draft(sid)
        parsed = ai.parse_model_output(
            '<patch anchor="1">FIRST</patch><patch anchor="3">THIRD</patch>'
        )
        events = await ai._land_turn(
            "tester", sid, parsed, save=True, autosave=False,
            selection="第二段正文",
        )
        self.assertEqual(self._texts(sid), ["第一段正文", "第二段正文", "第三段正文"])
        self.assertTrue(any(e.get("type") == "warn" for e in events), events)

    async def test_partial_out_of_scope_keeps_only_selected(self):
        """只有落在选区里的那条留下；越界的丢掉，并明确提示。"""
        sid = "guard-partial"
        self._draft(sid)
        parsed = ai.parse_model_output(
            '<patch anchor="1">FIRST</patch><patch anchor="2">SECOND-EN</patch>'
        )
        events = await ai._land_turn(
            "tester", sid, parsed, save=True, autosave=False,
            selection="第二段正文",
        )
        self.assertEqual(self._texts(sid), ["第一段正文", "SECOND-EN", "第三段正文"])
        self.assertTrue(any(e.get("type") == "warn" for e in events), events)

    async def test_final_doc_is_blocked_in_scoped_turn(self):
        """选区回合里整篇重写必须拦下（用户从没要过整篇）。"""
        sid = "guard-final-doc"
        self._draft(sid)
        parsed = ai.parse_model_output(
            "<final-doc>## 一\n\nEN1\n\n## 二\n\nEN2\n\n## 三\n\nEN3\n</final-doc>"
        )
        events = await ai._land_turn(
            "tester", sid, parsed, save=True, autosave=False,
            selection="第二段正文",
        )
        self.assertEqual(self._texts(sid), ["第一段正文", "第二段正文", "第三段正文"])
        self.assertTrue(any(e.get("type") == "warn" for e in events), events)

    async def test_style_all_is_not_respected_in_scoped_turn(self):
        """选区回合里模型的 anchor="all" 不予尊重：样式只落在选中的那段文字上。"""
        sid = "guard-style-all"
        self._draft(sid)
        parsed = ai.parse_model_output('<doc-style anchor="all" body-size="16"/>')
        events = await ai._land_turn(
            "tester", sid, parsed, save=True, autosave=False,
            selection="第二段正文",
        )
        # 进整篇 styles 的话，下次整篇导出会把选区字号铺到全文——换个方式复发原 bug
        self.assertEqual(ai.load_draft("tester", sid).get("styles") or {}, {})
        styles = [e for e in events if e.get("type") == "style"]
        self.assertEqual(len(styles), 1, events)
        # 作用域标识是 "selection"（用选区文字定位段落），不是草稿段号
        self.assertEqual(styles[0].get("anchor"), "selection")
        self.assertEqual(styles[0].get("scope"), "selection")

    async def test_style_never_falls_back_to_whole_doc(self):
        """即使选区与草稿对不上，样式也绝不落到整篇 styles（那正是全篇放大的成因）。"""
        sid = "guard-style-miss"
        self._draft(sid)
        parsed = ai.parse_model_output('<doc-style body-size="16"/>')
        events = await ai._land_turn(
            "tester", sid, parsed, save=True, autosave=False,
            selection="草稿里根本没有的这句话",
        )
        self.assertEqual(ai.load_draft("tester", sid).get("styles") or {}, {})
        styles = [e for e in events if e.get("type") == "style"]
        self.assertEqual(len(styles), 1, events)
        self.assertEqual(styles[0].get("anchor"), "selection")

    async def test_unscoped_turn_keeps_old_behaviour(self):
        """没有选区时一切照旧：模型要整篇就是整篇（回归保护）。"""
        sid = "guard-no-selection"
        self._draft(sid)
        parsed = ai.parse_model_output('<doc-style body-size="16"/>')
        events = await ai._land_turn("tester", sid, parsed, save=True, autosave=False)
        styles = [e for e in events if e.get("type") == "style"]
        self.assertEqual(len(styles), 1, events)
        self.assertEqual(styles[0].get("styles"), {"size_pt": 16.0})
        self.assertEqual(ai.load_draft("tester", sid).get("styles"), {"size_pt": 16.0})


class SelectionReplacementTest(unittest.IsolatedAsyncioTestCase):
    """文字级替换：用「选中的文字」当坐标，不再依赖段落编号。

    这条路是「草稿整篇只有一块、选区作用域退化成全篇」的正解：用户已经告诉了
    我们他选了哪段文字，直接在文档里换掉它就行——不需要知道那是第几段。
    """

    def test_parse_replace_selection(self):
        out = ai.parse_model_output("<replace-selection>新的文字</replace-selection>")
        self.assertEqual(out["replace_selection"], "新的文字")
        # 协议标记绝不能漏进正文
        self.assertNotIn("replace-selection", out["prose"])

    def test_locate_span_ignores_whitespace(self):
        """编辑器复制出来的选区带换行/缩进：按去空白对齐也要能定位回原文区间。"""
        text = "今夜，我坐在屏幕前，看着布宜诺斯艾利斯。"
        span = docx_revise.locate_selection_span(text, "我坐在屏幕前，\n看着布宜诺斯")
        self.assertIsNotNone(span)
        start, end = span
        self.assertEqual(text[start:end], "我坐在屏幕前，看着布宜诺斯")

    def test_replace_keeps_surrounding_text(self):
        merged = docx_revise.replace_selection_in_text(
            "前半句。中间这句要被换掉。后半句。", "中间这句要被换掉。", "NEW."
        )
        self.assertEqual(merged, "前半句。NEW.后半句。")

    def test_replace_returns_none_when_not_found(self):
        self.assertIsNone(
            docx_revise.replace_selection_in_text("原文", "不存在的片段", "X")
        )

    def test_whole_paragraph_paste_is_not_duplicated(self):
        """模型把整段塞进标签时，不能把整段插进选区位置（那样文字会重复）。"""
        text = "第一句。第二句。第三句。"
        # 只要求改样式时模型常把整段原样交回：结果应当原样，而不是重复一遍
        self.assertEqual(
            docx_revise.replace_selection_in_text(text, "第二句。", text), text
        )
        # 原样交回「选中的那几句」→ 这是片段，按片段处理，结果同样不变
        self.assertEqual(
            docx_revise.replace_selection_in_text(text, "第二句。", "第二句。"), text
        )

    def test_appending_keeps_the_prefix(self):
        """在选区后面追加内容时，绝不能把段落前半截吞掉。

        「原文 + 追加」这种形态很容易被误判成整段（因为新文本里带着原文），
        判成整段就会把选区之前的文字一起换掉。判据必须看**选区之外**的段落文字。
        """
        text = "第一句。第二句。第三句。"
        merged = docx_revise.replace_selection_in_text(
            text, "第二句。", "第二句。又加了一句。"
        )
        self.assertEqual(merged, "第一句。第二句。又加了一句。第三句。")

    def test_blocks_sync_only_touches_hit_block(self):
        blocks = [
            {"heading": "一", "text": "第一段正文"},
            {"heading": "二", "text": "第二段正文"},
        ]
        out = ai._replace_selection_in_blocks(blocks, "第二段正文", "SECOND")
        self.assertEqual([b["text"] for b in out], ["第一段正文", "SECOND"])
        self.assertIsNone(ai._replace_selection_in_blocks(blocks, "没有这句", "X"))

    def test_blocks_sync_treats_unchanged_as_hit(self):
        """原样返回也算命中——否则「只要求改样式」那一轮会误报「没能定位」。"""
        blocks = [{"heading": "一", "text": "第一句。第二句。"}]
        out = ai._replace_selection_in_blocks(blocks, "第二句。", "第二句。")
        self.assertIsNotNone(out)
        self.assertEqual([b["text"] for b in out], ["第一句。第二句。"])

    async def test_land_turn_replaces_only_selection(self):
        """落文：只有选中的那半句被换掉，段内其余文字原样。"""
        sid = "repl-turn-01"
        ai.save_draft("tester", sid, {
            "title": "测试稿",
            "blocks": [{"heading": "一", "text": "第一段正文，讲的是A方案。"}],
        })
        parsed = ai.parse_model_output(
            "<replace-selection>讲的是B方案。</replace-selection>"
        )
        events = await ai._land_turn(
            "tester", sid, parsed, save=True, autosave=False,
            selection="讲的是A方案。",
        )
        self.assertEqual(
            ai.load_draft("tester", sid)["blocks"][0]["text"],
            "第一段正文，讲的是B方案。",
        )
        self.assertTrue(any(e.get("type") == "draft" for e in events), events)

    async def test_land_turn_ignores_paragraph_patches(self):
        """模型既给替换、又给段落 patch：只认替换（否则等于变相放宽了范围）。"""
        sid = "repl-turn-02"
        ai.save_draft("tester", sid, {
            "title": "测试稿",
            "blocks": [
                {"heading": "一", "text": "甲。"},
                {"heading": "二", "text": "乙。"},
            ],
        })
        parsed = ai.parse_model_output(
            "<replace-selection>乙改。</replace-selection>"
            '<patch anchor="1">甲也改。</patch>'
        )
        events = await ai._land_turn(
            "tester", sid, parsed, save=True, autosave=False,
            selection="乙。",
        )
        self.assertEqual(
            [b["text"] for b in ai.load_draft("tester", sid)["blocks"]],
            ["甲。", "乙改。"],
        )
        self.assertTrue(
            any(
                e.get("type") == "warn" and "忽略" in str(e.get("message"))
                for e in events
            ),
            events,
        )

    async def test_real_docx_replace_keeps_font_and_no_revision(self):
        """真 docx 上验证三件事：文字换对了、其余段没动、**新文字继承了字体**。

        字体这条是「韩文/中文变成 □」的防线：新 run 必须带上原段的 `rPr`
        （东亚字体就在里面），否则 Word / Collabora 会拿西文字体去排这些字符。
        同时确认**没有留下修订**——改动是直接落地的，不再有叠加态。
        """
        from docx import Document
        import io

        if ai._skill_tool("md_to_docx.py", "") is None:
            self.skipTest("技能包未就绪（md_to_docx.py 不在 SKILLS_DIR），跳过转换测试")
        md = "## 一\n\n第一段正文，讲的是A方案。\n\n## 二\n\n第二段正文。\n"
        with tempfile.TemporaryDirectory() as tmp:
            base = await ai._convert_docx(
                tool="md_to_docx.py", src=md.encode("utf-8"),
                src_suffix=".md", out_suffix=".docx", root=Path(tmp), prefer="",
            )
        data, applied = docx_revise.apply_selection_replacement(
            base, "讲的是A方案。", "한국어 문장입니다."
        )
        self.assertEqual(applied, 1)
        doc = Document(io.BytesIO(data))
        paras = [p for p in doc.paragraphs if p.text.strip()]
        target = [p for p in paras if "한국어" in p.text][0]
        # 换对了：段内其余文字保留，只有选中的那句被替换
        self.assertEqual(target.text, "第一段正文，한국어 문장입니다.")
        # 直接落地：不留 w:ins / w:del（叠加态是这一整类不稳定的源头）
        self.assertEqual(list(target._element.iter(docx_revise._q("w:ins"))), [])
        self.assertEqual(list(target._element.iter(docx_revise._q("w:del"))), [])
        # 字体继承：新文字必须带着 rFonts（东亚字体），否则整片方框
        run = [r for r in target.runs if "한국어" in r.text][0]
        rPr = run._element.find(docx_revise._q("w:rPr"))
        self.assertIsNotNone(rPr, "新插入的文字没有 rPr——字体丢了")
        rfonts = rPr.find(docx_revise._q("w:rFonts"))
        self.assertIsNotNone(rfonts, "rPr 里没有 rFonts——中文/韩文会渲染成方框")
        self.assertEqual(rfonts.get(docx_revise._q("w:eastAsia")), "宋体")
        # 没被选中的段落一字不动
        other = [p for p in paras if "第二段正文" in p.text][0]
        self.assertEqual(other.text, "第二段正文。")

    async def test_markdown_bold_becomes_real_bold(self):
        """`**加粗**` 要变成真加粗，绝不能让 `**` 变成文档里的字面符号。"""
        from docx import Document
        import io

        if ai._skill_tool("md_to_docx.py", "") is None:
            self.skipTest("技能包未就绪（md_to_docx.py 不在 SKILLS_DIR），跳过转换测试")
        md = "## 一\n\n第一段正文，讲的是A方案。\n"
        with tempfile.TemporaryDirectory() as tmp:
            base = await ai._convert_docx(
                tool="md_to_docx.py", src=md.encode("utf-8"),
                src_suffix=".md", out_suffix=".docx", root=Path(tmp), prefer="",
            )
        data, applied = docx_revise.apply_selection_replacement(
            base, "讲的是A方案。", "**加粗英文 BOLD**"
        )
        self.assertEqual(applied, 1)
        doc = Document(io.BytesIO(data))
        target = [p for p in doc.paragraphs if "加粗英文" in p.text][0]
        self.assertNotIn("**", target.text, "markdown 标记被当成字面符号写进了文档")
        run = [r for r in target.runs if "加粗英文" in r.text][0]
        rPr = run._element.find(docx_revise._q("w:rPr"))
        self.assertIsNotNone(rPr.find(docx_revise._q("w:b")), "加粗标记没变成真加粗")

    async def test_append_text_keeps_the_original(self):
        """「在这段后面加一句」：原文保留、新内容接在后面。"""
        from docx import Document
        import io

        if ai._skill_tool("md_to_docx.py", "") is None:
            self.skipTest("技能包未就绪（md_to_docx.py 不在 SKILLS_DIR），跳过转换测试")
        md = "## 一\n\n第一段正文，讲的是A方案。\n"
        with tempfile.TemporaryDirectory() as tmp:
            base = await ai._convert_docx(
                tool="md_to_docx.py", src=md.encode("utf-8"),
                src_suffix=".md", out_suffix=".docx", root=Path(tmp), prefer="",
            )
        data, applied = docx_revise.apply_selection_replacement(
            base, "讲的是A方案。", "讲的是A方案。另外补充一句。"
        )
        self.assertEqual(applied, 1)
        doc = Document(io.BytesIO(data))
        target = [p for p in doc.paragraphs if "另外补充一句" in p.text][0]
        self.assertEqual(target.text, "第一段正文，讲的是A方案。另外补充一句。")

    async def test_locates_text_marked_as_deleted(self):
        """文档里若**已经有**修订（历史文档 / 用户在编辑器里改过），按「已被标记
        删除」的原文也要定位得到——否则在同一段上再改一次就会失败。"""
        from docx import Document
        import io

        if ai._skill_tool("md_to_docx.py", "") is None:
            self.skipTest("技能包未就绪（md_to_docx.py 不在 SKILLS_DIR），跳过转换测试")
        md = "## 一\n\n第一段正文，讲的是A方案。\n"
        with tempfile.TemporaryDirectory() as tmp:
            base = await ai._convert_docx(
                tool="md_to_docx.py", src=md.encode("utf-8"),
                src_suffix=".md", out_suffix=".docx", root=Path(tmp), prefer="",
            )
        # 手动造出「原文已被标记删除」的段落，模拟带修订的历史文档
        doc = Document(io.BytesIO(base))
        para = [p for p in doc.paragraphs if "讲的是A方案" in p.text][0]
        rev = docx_revise._Revisions()
        for run in list(para._element.findall(docx_revise._q("w:r"))):
            docx_revise._wrap(para._element, run, "del", rev)
        buf = io.BytesIO()
        doc.save(buf)
        again = Document(io.BytesIO(buf.getvalue()))
        self.assertIsNotNone(
            docx_revise._find_paragraph(again, "讲的是A方案。"),
            "被修订删除的原文定位不到——同一段再改一次就会失败",
        )


class DraftMarkdownNormalizeTest(unittest.IsolatedAsyncioTestCase):
    """`__粗体__` 必须统一成 `**粗体**`，否则会以字面符号写进文档。

    症状：重新打开文档时每段前面多出一对 `__`（回读带回来一层 → 导出写进文档）。
    成因：导出用的 `md_to_docx.py` 只认 `**…**`，`__…__` 会被当普通字符。
    """

    def test_underscore_bold_becomes_star(self):
        self.assertEqual(
            ai._normalize_inline_markdown("__今夜，我坐在屏幕前。__"),
            "**今夜，我坐在屏幕前。**",
        )

    def test_single_underscore_untouched(self):
        """单个下划线不碰——免得误伤正文里的普通符号。"""
        self.assertEqual(ai._normalize_inline_markdown("a_b_c"), "a_b_c")

    def test_draft_to_markdown_normalizes(self):
        draft = {"blocks": [{"heading": "一", "text": "__整段都被加粗了__"}]}
        md = ai.draft_to_markdown(draft)
        self.assertIn("**整段都被加粗了**", md)
        self.assertNotIn("__", md)

    async def test_exported_docx_has_no_literal_underscores(self):
        """真导出一次：docx 里不能出现 `__` 字面，而且加粗要真的生效。"""
        from docx import Document
        import io

        if ai._skill_tool("md_to_docx.py", "") is None:
            self.skipTest("技能包未就绪（md_to_docx.py 不在 SKILLS_DIR），跳过转换测试")
        md = ai.draft_to_markdown({"blocks": [{"heading": "一", "text": "__整段都被加粗了__"}]})
        with tempfile.TemporaryDirectory() as tmp:
            data = await ai._convert_docx(
                tool="md_to_docx.py", src=md.encode("utf-8"),
                src_suffix=".md", out_suffix=".docx", root=Path(tmp), prefer="",
            )
        doc = Document(io.BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertNotIn("__", text, "`__` 被当成字面符号写进了文档")
        bold_runs = [r for p in doc.paragraphs for r in p.runs if r.font.bold]
        self.assertTrue(bold_runs, "加粗标记没转换成真加粗")


if __name__ == "__main__":
    unittest.main()
